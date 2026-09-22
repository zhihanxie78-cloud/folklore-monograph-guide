# -*- coding: utf-8 -*-
"""Markdown → Word 转换器（book-summary-skill 配套工具）

用法:
    1. 复制本文件到 {书名}_详细梳理/ 文件夹
    2. 修改下方 BASE_DIR 与 FILES 两个变量
    3. cd "{书名}_详细梳理" && PYTHONUTF8=1 python md_to_docx.py

已内置:
    - 微软雅黑字体（中英文）、Table Grid 表格、h1 居中
    - 粗体/斜体/行内代码、引用块(▎灰字)、代码块(Consolas 9pt)
    - 有序/无序列表、水平线
依赖: pip install python-docx
"""
import os, re, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

# ============ 修改这两处 ============
BASE_DIR = r"D:/路径/{书名}_详细梳理"
FILES = [
    ("README.md", "README.docx"),
    ("01_知识图谱/知识图谱.md", "01_知识图谱/知识图谱.docx"),
    ("02_结构化提取/结构化提取.md", "02_结构化提取/结构化提取.docx"),
    ("03_学习笔记/学习笔记与资源索引.md", "03_学习笔记/学习笔记与资源索引.docx"),
]
# ====================================


def add_run_with_font(p, text, bold=False, italic=False, code=False, color=None, size=None):
    run = p.add_run(text)
    if bold: run.bold = True
    if italic: run.italic = True
    if code:
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
    elif size:
        run.font.size = Pt(size)
    if color: run.font.color.rgb = color
    return run


def parse_inline(text, p):
    pattern = r'(\*\*(.+?)\*\*|(?<!\*)\*(.+?)\*(?!\*)|`(.+?)`)'
    last_end = 0
    for m in re.finditer(pattern, text):
        if m.start() > last_end:
            p.add_run(text[last_end:m.start()])
        if m.group(2): add_run_with_font(p, m.group(2), bold=True)
        elif m.group(3): add_run_with_font(p, m.group(3), italic=True)
        elif m.group(4): add_run_with_font(p, m.group(4), code=True)
        last_end = m.end()
    if last_end < len(text):
        p.add_run(text[last_end:])


def parse_table_line(line):
    line = line.strip()
    if line.startswith('|'): line = line[1:]
    if line.endswith('|'): line = line[:-1]
    return [c.strip() for c in line.split('|')]


def is_separator_line(line):
    return bool(re.match(r'^[\|\s\-:]+$', line.strip()))


def set_cell_text(cell, text, bold=False):
    cell.text = ''
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.font.size = Pt(9)
    run.font.name = '微软雅黑'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run.bold = bold


def convert_md_to_docx(md_rel_path, docx_rel_path):
    md_path = os.path.join(BASE_DIR, md_rel_path)
    docx_path = os.path.join(BASE_DIR, docx_rel_path)
    print(f"Converting: {md_rel_path}")
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = '微软雅黑'
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
    lines = content.split('\n'); i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip(): i += 1; continue
        if line.startswith('```'):
            code_lines = []; i += 1
            while i < len(lines) and not lines[i].rstrip().startswith('```'):
                code_lines.append(lines[i]); i += 1
            if code_lines:
                p = doc.add_paragraph()
                run = p.add_run('\n'.join(code_lines))
                run.font.name = 'Consolas'; run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                p.paragraph_format.left_indent = Cm(1)
            i += 1; continue
        if '|' in line and line.strip().startswith('|'):
            table_lines = [line]; j = i + 1
            while j < len(lines) and '|' in lines[j] and lines[j].strip().startswith('|'):
                table_lines.append(lines[j].rstrip()); j += 1
            data_lines = [l for l in table_lines if not is_separator_line(l)]
            if data_lines:
                cols = len(parse_table_line(data_lines[0]))
                parsed_rows = []
                for tl in data_lines:
                    cells = parse_table_line(tl)
                    while len(cells) < cols: cells.append('')
                    parsed_rows.append(cells[:cols])
                table = doc.add_table(rows=len(parsed_rows), cols=cols)
                table.style = 'Table Grid'; table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for row_idx, cells in enumerate(parsed_rows):
                    for col_idx, cell_text in enumerate(cells):
                        set_cell_text(table.cell(row_idx, col_idx), cell_text, bold=(row_idx == 0))
                doc.add_paragraph()
            i = j; continue
        h_match = re.match(r'^(#{1,6})\s+(.+)$', line)
        if h_match:
            level = len(h_match.group(1))
            text = re.sub(r'[*`]', '', h_match.group(2))
            h = doc.add_heading(text, level=level)
            if level == 1: h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1; continue
        if re.match(r'^[-_*]{3,}\s*$', line):
            p = doc.add_paragraph('─' * 60); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1; continue
        if line.startswith('>'):
            text = re.sub(r'^>\s*', '', line)
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(1)
            prefix = '▎ '; p.add_run(prefix); parse_inline(text, p)
            for run in p.runs:
                if run.text != prefix: run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1; continue
        ul_match = re.match(r'^(\s*)[\-\*]\s+(.+)$', line)
        if ul_match:
            indent = len(ul_match.group(1)); text = ul_match.group(2)
            p = doc.add_paragraph(); p.style = doc.styles['List Bullet']
            for run in p.runs: run.text = ''
            parse_inline(text, p)
            if indent > 0: p.paragraph_format.left_indent = Cm(1.5 + indent * 0.5)
            i += 1; continue
        ol_match = re.match(r'^(\s*)\d+[\.\)]\s+(.+)$', line)
        if ol_match:
            text = ol_match.group(2)
            p = doc.add_paragraph(); p.style = doc.styles['List Number']
            for run in p.runs: run.text = ''
            parse_inline(text, p)
            i += 1; continue
        p = doc.add_paragraph(); parse_inline(line, p); i += 1
    try:
        doc.save(docx_path); print(f"  -> Saved: {docx_rel_path}")
    except Exception as e:
        print(f"  -> Save error: {e}")


if __name__ == "__main__":
    success = 0
    for md_rel, docx_rel in FILES:
        md_full = os.path.join(BASE_DIR, md_rel)
        if os.path.exists(md_full):
            try:
                convert_md_to_docx(md_rel, docx_rel); success += 1
            except Exception as e:
                import traceback
                print(f"Error: {md_rel}: {e}"); traceback.print_exc()
        else:
            print(f"File not found: {md_full}")
    print(f"\nDone! {success}/{len(FILES)} files converted.")
